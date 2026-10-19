"""제안서 작성기.

입력: BidRequirement + 후보 판정(Evaluation) + 회사 프로필(company_profile.json) + (선택) 비용 분석
출력: Proposal(섹션·요소 트리) → HTML(인쇄용 A4) / DOCX

정직성 규칙
  - 확인되지 않은 값은 지어내지 않고 '[확인 필요: …]' 로 남긴다 (화면·문서에서 강조).
  - 외부 제출본(external)에는 원가·마진·내부 판정 근거를 넣지 않는다. 부록(내부용)은 internal 모드에서만.
  - LLM 다듬기는 선택이며, 결과 문장에 원문에 없던 숫자가 생기면 폐기하고 원문을 쓴다.
"""
from __future__ import annotations

import html
import json
import re
from datetime import date
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field

from b2b.costing import CostAnalysis
from b2b.models import BidRequirement, Evaluation
from shared.config import get_settings
from shared.llm import LLMClient

Mode = Literal["external", "internal"]
TODO = re.compile(r"\[확인 필요[^\]]*\]")


class Element(BaseModel):
    kind: Literal["para", "bullets", "table", "callout"]
    text: str = ""
    items: list[str] = Field(default_factory=list)
    headers: list[str] = Field(default_factory=list)
    rows: list[list[str]] = Field(default_factory=list)
    tone: str = "info"


class Section(BaseModel):
    id: str
    title: str
    elements: list[Element]
    internal_only: bool = False


class Proposal(BaseModel):
    title: str
    subtitle: str
    bid_id: str | None
    agency: str | None
    company: str
    created: date
    mode: Mode
    sections: list[Section]
    todo_count: int = 0


def won(v: int | None) -> str:
    return "[확인 필요: 금액]" if v is None else f"{v:,}원"


def load_profile(path: str | None = None) -> dict:
    p = Path(path or get_settings().company_profile_path)
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _todo(v: Any, label: str) -> str:
    if v in (None, "", []) or (isinstance(v, str) and v.startswith("[")):
        return f"[확인 필요: {label}]"
    return str(v)


def build(req: BidRequirement, evals: list[Evaluation], *, profile: dict | None = None,
          cost: CostAnalysis | None = None, mode: Mode = "internal", today: date | None = None,
          polish: bool = False, source_status: str | None = None) -> Proposal:
    """source_status: 공고 출처가 확인되지 않았을 때 그 상태 이름 (예: '원문 대조 전'). 확인된 공고면 None."""
    pf = profile if profile is not None else load_profile()
    today = today or date.today()
    company = _todo(pf.get("company_name"), "제안사명")
    agency = req.agency or "[확인 필요: 발주기관]"
    title = req.title or f"{agency} 임직원 선물 공급 사업"
    n_items = sum(c.count for c in req.categories) or None
    cat_txt = "·".join(f"{c.name} {c.count}종" for c in req.categories) or "[확인 필요: 품목 구성]"
    total = req.computed_total()
    vat = "VAT 포함" if req.vat_included else "VAT 별도" if req.vat_included is False else "[확인 필요: VAT 기준]"
    d_ship = f"{req.delivery_business_days}영업일" if req.delivery_business_days else "[확인 필요: 배송 기한]"
    d_ex = f"{req.exchange_business_days}영업일" if req.exchange_business_days else "[확인 필요: 교환 기한]"
    d_setup = f"{req.setup_days}일" if req.setup_days else "[확인 필요: 구축 기한]"

    S: list[Section] = []

    # 1. 제안 요약
    summary = [
        f"{agency} 임직원 {req.headcount:,}명을 대상으로 {cat_txt} 중 1인 1개를 직접 선택하는 폐쇄몰형 선물 공급을 제안합니다."
        if req.headcount else f"{agency} 임직원 대상 {cat_txt} 선택형 선물 공급을 제안합니다.",
        f"계약 후 {d_setup} 이내 폐쇄몰을 열고, 선택일로부터 {d_ship} 이내 개별 주소 배송, 교환은 접수 후 {d_ex} 이내 완료를 목표로 운영합니다.",
        f"1인당 {won(req.budget_per_person)}({vat}) 단가로 제안하며, " +
        ("실제 지급수량 기준으로 정산합니다." if req.settlement else "정산 방식은 [확인 필요: 정산 기준]에 따릅니다."),
    ]
    if req.min_online_price:
        summary.append(f"모든 상품은 온라인 기준 판매가 {req.min_online_price:,}원 이상 상품으로 구성하고, 동일 옵션의 온라인 판매가 증빙을 함께 제출합니다.")
    if polish:
        summary = _polish(summary)
    S.append(Section(id="summary", title="1. 제안 요약", elements=[Element(kind="bullets", items=summary)]))

    # 2. 사업 이해
    src_warn = ([Element(kind="callout", tone="warn",
                         text=f"[확인 필요: 공고 원문 대조 — 현재 '{source_status}' 상태의 공고 정보로 작성됨]")]
                if source_status else [])
    S.append(Section(id="understanding", title="2. 사업 이해", elements=src_warn + [
        Element(kind="para", text=f"본 사업은 {agency} 임직원에게 선택권을 보장하는 선물 지급 사업으로, 짧은 구축 기한과 "
                                  f"전국 개별 배송, 교환 대응 속도가 품질을 좌우합니다. 제안사는 상품 다양성보다 "
                                  f"'약속한 기한 안에 정확한 상품이 도착하는 것'을 최우선 목표로 둡니다."),
        Element(kind="table", headers=["항목", "내용"], rows=[
            ["사업명", title], ["발주기관", agency], ["공고번호", req.bid_id or "[확인 필요: 공고번호]"],
            ["예정 인원", f"{req.headcount:,}명" if req.headcount else "[확인 필요: 인원]"],
            ["1인당 단가", f"{won(req.budget_per_person)} ({vat})"],
            ["예정 총액", f"{won(total)} (실제 지급수량 정산)" if total else "[확인 필요: 총액]"],
            ["상품 구성", f"{cat_txt}" + (f" (총 {n_items}종)" if n_items else "")],
            ["상품 가치 조건", f"온라인 기준 판매가 {req.min_online_price:,}원 이상" if req.min_online_price else "[확인 필요]"],
        ]),
    ]))

    # 3. 요구사항 대응표
    rows = [
        ["폐쇄몰 구축·상품 등재", f"계약일로부터 {d_setup} 이내", "D+0~D+5 구축 일정(5장) · 임직원 인증·1인 1선택", "대응"],
        ["상품 구성", cat_txt, "카테고리별 후보 복수 확보 후 공급사 확약분만 등재(4장)", "후보 확인 중"],
        ["온라인가 하한", f"{req.min_online_price:,}원 이상" if req.min_online_price else "-",
         "동일 옵션 온라인 판매 경로·가격 캡처 제출", "증빙 준비 중"],
        ["배송", f"선택일부터 {d_ship} 이내 국내 지정 주소", "공급사별 출고 마감·송장 회신 기준 계약 반영", "공급사 확약 필요"],
        ["교환", f"접수일부터 {d_ex} 이내", "교환 접수 → 회수·재발송 동시 진행, 비용 부담 기준 사전 합의", "공급사 확약 필요"],
        ["정산", req.settlement or "-", "실제 지급 수량 확정 후 정산 자료 제출", "대응"],
    ]
    S.append(Section(id="compliance", title="3. 요구사항 대응표", elements=[
        Element(kind="table", headers=["요구사항", "공고 조건", "제안 내용", "상태"], rows=rows),
        Element(kind="callout", tone="info", text="'공급사 확약 필요' 항목은 견적·확약서 회신 후 '대응'으로 바꿔 제출합니다."),
    ]))

    # 4. 상품 구성안
    els: list[Element] = [Element(kind="para", text="카테고리별로 조건을 충족할 수 있는 후보를 복수로 확보한 뒤, 공급사 견적·배송 확약·"
                                                     "온라인가 증빙이 갖춰진 상품만 최종 구성에 넣습니다.")]
    for need in req.categories or []:
        cands = [e for e in evals if e.candidate.category == need.name and e.decision != "EXCLUDED"
                 and not e.candidate.is_sample]
        if mode == "external" and req.min_online_price:  # 제출본: 하한 이상으로 관측된 후보만
            cands = [e for e in cands if (e.candidate.observed_price or 0) >= req.min_online_price]
        trows = []
        for e in cands[: max(need.count + 2, 3)]:
            c = e.candidate
            price = f"{c.observed_price:,}원 (관측 {c.observed_date})" if c.observed_price else "[확인 필요: 온라인가]"
            if e.decision == "ELIGIBLE":
                status = "확정"
            elif mode == "external":
                status = "[확인 필요: 공급사 확약 후 확정]"
            else:
                status = f"[확인 필요: {', '.join(x.label for x in e.unknown[:3])}]"
            row = [c.name, c.option or "[확인 필요: 옵션]", price, status]
            if mode == "internal":
                row.append(c.supplier)
            trows.append(row)
        if not trows:
            trows.append([f"[확인 필요: {need.name} 후보 발굴]", "-", "-", "후보 없음"] + (["-"] if mode == "internal" else []))
        hdr = ["상품", "옵션·구성", "온라인 판매가", "상태"] + (["공급 경로"] if mode == "internal" else [])
        els += [Element(kind="para", text=f"■ {need.name} {need.count}종"), Element(kind="table", headers=hdr, rows=trows)]
    els.append(Element(kind="callout", tone="warn",
                       text="온라인 판매가는 관측 시점 값이며 시장 최저가 검증 전입니다. 제출 전 동일 옵션 복수 판매경로 가격을 대조합니다."))
    S.append(Section(id="products", title="4. 상품 구성안", elements=els))

    # 5. 폐쇄몰
    cap = pf.get("capabilities", {})
    S.append(Section(id="mall", title="5. 폐쇄몰 구축·운영", elements=[
        Element(kind="para", text=f"구축 방식: {_todo(cap.get('closed_mall'), '폐쇄몰 플랫폼·구축 방식')}"),
        Element(kind="table", headers=["일정", "작업", "산출물"], rows=[
            ["D+0", "계약·임직원 명단/배송 정책 수령, 상품 상세정보 확보", "착수 회의록"],
            ["D+1", "상품 등재, 임직원 인증(사번·휴대폰) 연동", "등재 목록"],
            ["D+2", "1인 1선택·배송지 입력·변경 마감 기능 점검", "기능 점검표"],
            ["D+3", "발주처 검수", "검수 의견"],
            ["D+4", "수정 반영, 안내 문자·공지 준비", "수정 내역"],
            [f"D+{req.setup_days or 5}", "오픈 및 선택 현황 대시보드 공유", "오픈 보고"],
        ]),
        Element(kind="bullets", items=["임직원 본인 인증 후 1인 1개 선택, 마감 전 변경 허용",
                                        "미선택자 독려 안내(기한·횟수는 발주처와 협의)",
                                        "발주처 담당자용 선택·배송 현황 실시간 조회"]),
    ]))

    # 6. 배송·교환·CS
    S.append(Section(id="delivery", title="6. 배송·교환·CS", elements=[
        Element(kind="table", headers=["단계", "기준", "관리 방법"], rows=[
            ["주문 전달", "선택 마감 당일", "공급사별 주문서 자동 분리·전달"],
            ["출고·배송", f"선택일부터 {d_ship} 이내 도착", "공급사 출고 마감 시각 계약 명시, 송장 회신 일 1회 집계"],
            ["지연 관리", "도착 기한 D-1 미출고 건", "발주처 사전 보고 및 대체 출고"],
            ["교환", f"접수일부터 {d_ex} 이내", "회수·재발송 병행, 비용 부담 기준 사전 합의"],
            ["CS", _todo(cap.get("cs"), "CS 운영 체계"), "문의 유형별 응대 기록, 주간 보고"],
        ]),
        Element(kind="para", text=f"물류: {_todo(cap.get('logistics'), '물류·배송 파트너')}"),
    ]))

    # 7. 가격 제안
    S.append(Section(id="price", title="7. 가격 제안", elements=[
        Element(kind="table", headers=["구분", "금액", "비고"], rows=[
            ["1인당 단가", won(req.budget_per_person), vat + ", 상품·포장·배송·폐쇄몰 운영 포함"],
            ["예정 총액", won(total), f"{req.headcount:,}명 기준, 실제 지급수량 정산" if req.headcount else "실제 지급수량 정산"],
        ]),
        Element(kind="para", text="도서산간 추가 배송비와 교환 비용은 단가에 포함하며 별도 청구하지 않는 것을 원칙으로 합니다. [확인 필요: 공급사 견적 확인 후 확정]"),
    ]))

    # 8. 품질·리스크
    S.append(Section(id="risk", title="8. 품질·리스크 관리", elements=[
        Element(kind="table", headers=["리스크", "대응"], rows=[
            ["품절·계절 공급 변동", "동급 이상 대체 상품을 사전 지정하고 발주처 승인 후 교체"],
            ["가격 변동", "견적 유효기간을 사업 기간 이상으로 확보, 만료 전 재확인"],
            ["신선식품 파손·품질", "출고 전 샘플 검수, 파손 시 우선 재발송"],
            ["배송 지연", "공급사별 출고 마감 관리, 지연 예상 건 사전 보고"],
            ["개인정보", "배송 정보는 배송 목적으로만 사용, 사업 종료 후 파기"],
        ]),
    ]))

    # 9. 회사 소개
    track = pf.get("track_record") or []
    S.append(Section(id="company", title="9. 제안사 소개", elements=[
        Element(kind="para", text=_todo(pf.get("intro"), "회사 소개")),
        Element(kind="table", headers=["항목", "내용"], rows=[
            ["회사명", company], ["대표자", _todo(pf.get("representative"), "대표자")],
            ["사업자등록번호", _todo(pf.get("business_no"), "사업자등록번호")], ["주소", _todo(pf.get("address"), "주소")],
            ["담당자", f"{_todo(pf.get('contact_name'), '담당자')} / {_todo(pf.get('contact_phone'), '연락처')} / {_todo(pf.get('contact_email'), '이메일')}"],
        ]),
        Element(kind="bullets", items=track or ["[확인 필요: 유사 사업 수행 실적 — 실제 실적만 기재]"]),
    ]))

    # ── 내부 부록 ──────────────────────────────────────────
    if mode == "internal":
        qs: dict[str, set[str]] = {}
        for e in evals:
            if e.decision == "EXCLUDED" or e.candidate.is_sample:
                continue
            for c in e.unknown:
                qs.setdefault(c.label, set()).add(e.candidate.name)
        S.append(Section(id="appendix_check", title="부록 A. 제출 전 확인 체크리스트 (내부용)", internal_only=True, elements=[
            Element(kind="table", headers=["확인 항목", "해당 후보"],
                    rows=[[k, ", ".join(sorted(v))[:160]] for k, v in qs.items()] or [["없음", "-"]]),
        ]))
        if cost:
            crow = [[l.label, won(l.amount) if l.amount is not None else "미확인", l.basis] for l in cost.lines]
            els2 = [Element(kind="table", headers=["비용 항목(1인)", "금액", "근거"], rows=crow),
                    Element(kind="callout", tone="warn",
                            text=("잔여액 " + won(cost.residual) + " (입력값 기준 — 견적 확정 전)") if cost.residual is not None
                            else f"미확인 {len(cost.unknown)}개 항목이 있어 잔여액(마진)을 확정하지 않습니다. 미확인 ≠ 0원.")]
            if cost.scenarios:
                els2.append(Element(kind="table", headers=["가정(공개가=매입가)", "가정 매입가", "추가비용 전 차액", "판단"],
                                    rows=[[s["candidate"], won(s["assumed_supply"]), won(s["gap_before_other_costs"]), s["verdict"]]
                                          for s in cost.scenarios]))
            els2.append(Element(kind="bullets", items=cost.notes))
            S.append(Section(id="appendix_cost", title="부록 B. 비용 구조 (내부용)", internal_only=True, elements=els2))

    prop = Proposal(title=title, subtitle="임직원 선물 공급 제안서" + (" — 내부 검토본" if mode == "internal" else " — 제출본 초안"),
                    bid_id=req.bid_id, agency=agency, company=company, created=today, mode=mode, sections=S)
    prop.todo_count = sum(len(TODO.findall(_flat(s))) for s in S if not s.internal_only)
    return prop


def _flat(s: Section) -> str:
    return " ".join([e.text for e in s.elements] + [x for e in s.elements for x in e.items] +
                    [c for e in s.elements for r in e.rows for c in r])


def _polish(lines: list[str]) -> list[str]:
    llm = LLMClient()
    if not llm.enabled:
        return lines
    sys = ("공공·기업 제안서 문체로 문장을 다듬어라. 주어진 사실만 사용하고 숫자·기한·금액을 바꾸거나 추가하지 마라. "
           "실적·인증·수상 등 새로운 사실을 만들지 마라. '[확인 필요: …]' 표시는 그대로 둬라. JSON {\"lines\": [...]} 만 출력.")
    try:
        out = llm.chat_json(sys, json.dumps({"lines": lines}, ensure_ascii=False)).get("lines", [])
    except Exception:
        return lines
    nums = lambda s: set(re.findall(r"\d[\d,]*", s))  # noqa: E731
    src = set().union(*(nums(x) for x in lines))
    if len(out) != len(lines) or any(not nums(x) <= src for x in out):
        return lines  # 새 숫자가 생기면 폐기
    return out


# ── 렌더러 ────────────────────────────────────────────────────
def _mark(s: str) -> str:
    return TODO.sub(lambda m: f'<mark>{m.group(0)}</mark>', html.escape(s))


def to_html(p: Proposal) -> str:
    body = []
    for s in p.sections:
        part = [f'<section class="{"internal" if s.internal_only else ""}"><h2>{html.escape(s.title)}</h2>']
        for e in s.elements:
            if e.kind == "para":
                part.append(f"<p>{_mark(e.text)}</p>")
            elif e.kind == "bullets":
                part.append("<ul>" + "".join(f"<li>{_mark(x)}</li>" for x in e.items) + "</ul>")
            elif e.kind == "callout":
                part.append(f'<div class="callout {e.tone}">{_mark(e.text)}</div>')
            elif e.kind == "table":
                part.append('<table><thead><tr>' + "".join(f"<th>{html.escape(h)}</th>" for h in e.headers) + "</tr></thead><tbody>"
                            + "".join("<tr>" + "".join(f"<td>{_mark(c)}</td>" for c in r) + "</tr>" for r in e.rows)
                            + "</tbody></table>")
        part.append("</section>")
        body.append("".join(part))
    todo = (f'<div class="todo">제출 전 채워야 할 <b>[확인 필요]</b> 항목 {p.todo_count}개 (노란 표시)</div>' if p.todo_count else "")
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(p.title)} 제안서</title>
<link rel="stylesheet" href="https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/variable/pretendardvariable.min.css">
<style>
@page{{size:A4;margin:18mm 16mm}}
body{{font-family:"Pretendard Variable",Pretendard,system-ui,sans-serif;color:#15181d;background:#f3f2ee;margin:0;line-height:1.7;font-size:14.5px}}
.doc{{max-width:820px;margin:24px auto;background:#fff;padding:56px 64px;box-shadow:0 2px 24px rgba(0,0,0,.06)}}
.cover{{border-bottom:3px solid #15181d;padding-bottom:28px;margin-bottom:28px}}
.cover .k{{font-size:13px;letter-spacing:.08em;color:#6b6f78}}
.cover h1{{font-size:30px;line-height:1.3;margin:10px 0 6px;letter-spacing:-.02em}}
.cover .s{{font-size:17px;color:#3b3f47}}
.cover dl{{display:grid;grid-template-columns:110px 1fr;gap:4px 12px;margin:22px 0 0;font-size:14px}}
.cover dt{{color:#6b6f78}}
h2{{font-size:19px;margin:34px 0 12px;padding-top:10px;border-top:1px solid #e3e1da;letter-spacing:-.01em}}
table{{border-collapse:collapse;width:100%;margin:10px 0 14px;font-size:13.5px}}
th,td{{border:1px solid #d9dde3;padding:8px 10px;text-align:left;vertical-align:top}}
th{{background:#eef1f4;font-weight:650}}
mark{{background:#fff1a8;padding:0 2px;border-radius:3px}}
.callout{{padding:10px 14px;border-radius:8px;margin:10px 0;font-size:13.5px}}
.callout.info{{background:#eef4fb;color:#1d3e66}} .callout.warn{{background:#fff4dc;color:#6b4700}}
section.internal{{background:#fafaf7;border:1px dashed #c8c4b8;padding:4px 18px 10px;margin-top:28px;border-radius:8px}}
.todo{{position:sticky;top:0;background:#15181d;color:#fff;padding:10px 16px;font-size:13.5px;text-align:center}}
@media print{{body{{background:#fff}}.doc{{box-shadow:none;margin:0;padding:0}}.todo{{display:none}}mark{{background:none;border-bottom:1.5px solid #d4a800}}}}
@media (max-width:700px){{.doc{{padding:28px 18px;margin:0}}}}
</style></head><body>{todo}<div class="doc">
<div class="cover"><div class="k">PROPOSAL{' · INTERNAL DRAFT' if p.mode == 'internal' else ''}</div>
<h1>{html.escape(p.title)}</h1><div class="s">{html.escape(p.subtitle)}</div>
<dl><dt>발주기관</dt><dd>{_mark(p.agency or '')}</dd><dt>공고번호</dt><dd>{_mark(p.bid_id or '[확인 필요: 공고번호]')}</dd>
<dt>제안사</dt><dd>{_mark(p.company)}</dd><dt>작성일</dt><dd>{p.created.isoformat()}</dd></dl></div>
{''.join(body)}</div></body></html>"""


def to_docx(p: Proposal, path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Malgun Gothic"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    st.font.size = Pt(10.5)

    def runs(par, text: str):
        pos = 0
        for m in TODO.finditer(text):
            if m.start() > pos:
                par.add_run(text[pos:m.start()])
            r = par.add_run(m.group(0))
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            pos = m.end()
        if pos < len(text):
            par.add_run(text[pos:])

    k = doc.add_paragraph()
    r = k.add_run("PROPOSAL" + (" · INTERNAL DRAFT" if p.mode == "internal" else ""))
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x6B, 0x6F, 0x78)
    doc.add_heading(p.title, level=0)
    doc.add_paragraph(p.subtitle)
    for label, val in [("발주기관", p.agency or ""), ("공고번호", p.bid_id or "[확인 필요: 공고번호]"),
                       ("제안사", p.company), ("작성일", p.created.isoformat())]:
        par = doc.add_paragraph()
        par.add_run(f"{label}  ").bold = True
        runs(par, val)
    if p.todo_count:
        doc.add_paragraph().add_run(f"※ 제출 전 [확인 필요] 항목 {p.todo_count}개를 채우고 이 줄을 삭제하세요.").italic = True
    doc.add_page_break()
    for s in p.sections:
        doc.add_heading(s.title, level=1)
        for e in s.elements:
            if e.kind in ("para", "callout"):
                runs(doc.add_paragraph(), ("※ " if e.kind == "callout" else "") + e.text)
            elif e.kind == "bullets":
                for x in e.items:
                    runs(doc.add_paragraph(style="List Bullet"), x)
            elif e.kind == "table":
                t = doc.add_table(rows=1, cols=len(e.headers))
                t.style = "Table Grid"
                for i, h in enumerate(e.headers):
                    t.rows[0].cells[i].text = ""
                    t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
                for row in e.rows:
                    cells = t.add_row().cells
                    for i, c in enumerate(row[: len(e.headers)]):
                        cells[i].text = ""
                        runs(cells[i].paragraphs[0], c)
                doc.add_paragraph()
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def to_pptx(p: Proposal, path: Path) -> Path:
    """제안서를 PowerPoint (PPTX)로 생성한다. 각 섹션을 슬라이드로 구성."""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 표지 슬라이드
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = p.title
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(0x15, 0x18, 0x1D)

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.7), Inches(0.6))
    sub.text_frame.paragraphs[0].text = p.subtitle
    sub.text_frame.paragraphs[0].font.size = Pt(18)
    sub.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x3B, 0x3F, 0x47)

    meta = slide.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.5))
    mf = meta.text_frame
    mf.word_wrap = True
    lines = [
        f"발주기관  {p.agency or '[확인 필요]'}",
        f"공고번호  {p.bid_id or '[확인 필요: 공고번호]'}",
        f"제안사    {p.company}",
        f"작성일    {p.created.isoformat()}",
    ]
    for i, line in enumerate(lines):
        para = mf.paragraphs[0] if i == 0 else mf.add_paragraph()
        para.text = line
        para.font.size = Pt(14)
        para.font.color.rgb = RGBColor(0x6B, 0x6F, 0x78)

    if p.todo_count:
        note = slide.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.4))
        note.text_frame.paragraphs[0].text = f"※ 제출 전 [확인 필요] 항목 {p.todo_count}개"
        note.text_frame.paragraphs[0].font.size = Pt(12)
        note.text_frame.paragraphs[0].font.italic = True
        note.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xB4, 0x7A, 0x00)

    # 섹션별 슬라이드
    for s in p.sections:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # 섹션 제목
        hbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.7))
        hp = hbox.text_frame.paragraphs[0]
        hp.text = s.title
        hp.font.size = Pt(24)
        hp.font.bold = True
        hp.font.color.rgb = RGBColor(0x15, 0x18, 0x1D)

        # 본문 영역
        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12.1), Inches(5.8))
        tf = body.text_frame
        tf.word_wrap = True
        first = True

        for e in s.elements:
            if e.kind == "para" or e.kind == "callout":
                para = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                prefix = "※ " if e.kind == "callout" else ""
                para.text = prefix + e.text
                para.font.size = Pt(14)
                para.space_after = Pt(8)
            elif e.kind == "bullets":
                for item in e.items:
                    para = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    para.text = "• " + item
                    para.font.size = Pt(14)
                    para.level = 0
                    para.space_after = Pt(4)
            elif e.kind == "table" and e.headers and e.rows:
                # 표는 텍스트로 간단히 표현 (PPTX 표는 복잡하므로 요약)
                para = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                para.text = " | ".join(e.headers)
                para.font.size = Pt(12)
                para.font.bold = True
                for row in e.rows[:8]:  # 슬라이드당 과도하지 않게
                    rp = tf.add_paragraph()
                    rp.text = " | ".join(str(c)[:40] for c in row[:len(e.headers)])
                    rp.font.size = Pt(11)
                    rp.space_after = Pt(2)
                if len(e.rows) > 8:
                    more = tf.add_paragraph()
                    more.text = f"... 외 {len(e.rows)-8}행 (상세는 Word/HTML 참고)"
                    more.font.size = Pt(10)
                    more.font.italic = True

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def to_hwpx(p: Proposal, path: Path) -> Path:
    """제안서를 HWPX로 생성한다. python-hwpx 사용 (한컴 오피스 설치 불필요)."""
    try:
        from hwpx import HwpxDocument
    except ImportError:
        raise RuntimeError("python-hwpx가 설치되지 않았습니다. pip install python-hwpx")

    doc = HwpxDocument.new()

    # 표지
    doc.add_heading(p.title, level=1)
    doc.add_paragraph(p.subtitle)
    doc.add_paragraph("")
    meta_lines = [
        f"발주기관: {p.agency or '[확인 필요]'}",
        f"공고번호: {p.bid_id or '[확인 필요: 공고번호]'}",
        f"제안사: {p.company}",
        f"작성일: {p.created.isoformat()}",
    ]
    for line in meta_lines:
        doc.add_paragraph(line)
    if p.todo_count:
        doc.add_paragraph(f"※ 제출 전 [확인 필요] 항목 {p.todo_count}개를 확인하세요.")
    doc.add_paragraph("")

    # 섹션
    for s in p.sections:
        doc.add_heading(s.title, level=2)
        for e in s.elements:
            if e.kind in ("para", "callout"):
                text = ("※ " if e.kind == "callout" else "") + e.text
                doc.add_paragraph(text)
            elif e.kind == "bullets":
                for item in e.items:
                    doc.add_paragraph("• " + item)
            elif e.kind == "table" and e.headers and e.rows:
                table = doc.add_table(rows=1 + len(e.rows), cols=len(e.headers))
                for i, h in enumerate(e.headers):
                    table.set_cell_text(0, i, h)
                for ri, row in enumerate(e.rows):
                    for ci, cell in enumerate(row[:len(e.headers)]):
                        table.set_cell_text(ri + 1, ci, str(cell))
                doc.add_paragraph("")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save_to_path(str(path))
    return path
